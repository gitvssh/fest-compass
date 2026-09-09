"""Build dated display boundaries from reviewed public sources, entirely offline.

Python 3.12; pyshp 2.3.1, pyproj 3.7.2, shapely 2.1.2, openpyxl 3.1.5.
Usage: python -m tools.build_region_boundaries --sources /path/to/downloads
The source directory must contain sgis-2025.zip and the two pinned XLSX files.
"""
import argparse
from collections import Counter
import hashlib
import io
import json
from pathlib import Path
import zipfile

import openpyxl
import pyproj
import shapefile
import shapely
from shapely.geometry import shape, mapping
from shapely.ops import transform, unary_union

from tools.audit_region_boundaries import ARCHIVE_SHA256, ROOT, digest_file

PINS = {
    "sgis-2025.zip": ARCHIVE_SHA256,
    "kssc-20250401.xlsx": "ff987cb2505a664c465131d30bb5138d012323b448b5ce42256cdd8811cc2f6e",
    "kssc-20260710.xlsx": "699d33c57828f8b2f7b2546360e6d65667f3d339d641f50ca757840a3f4b4a4c",
}
SOURCE = "https://www.data.go.kr/data/15129688/fileData.do"
VERSION = "sgis-20250630-kssc-20260710-v1"


def read_links(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    for r in book["2-2. 연계표_행정동 및 법정동(기준시점)"].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[5], (int, float)) or not isinstance(r[8], (int, float)):
            continue
        code = str(int(r[5]))
        if len(code) not in (2, 5):
            continue
        item = {"province": r[0], "name": r[2], "legal": str(int(r[8]))}
        if code in result and result[code] != item:
            raise ValueError("Ambiguous official classification")
        result[code] = item
    book.close()
    return result


def link_region(row, old, current, available):
    # KTO's unusual Sejong pair is preserved; exact names and official legal
    # code are still required. Numeric equality between code systems is unsafe.
    key = ("36110" if (row["provinceCode"], row["districtCode"]) == ("36110", "36110")
           else row["provinceCode"] + row["districtCode"])
    candidates = [code for code, item in current.items() if len(code) == 5
                  and item == {"province": row["provinceName"], "name": row["districtName"],
                               "legal": key + "00000"}]
    if len(candidates) != 1:
        return {"status": "unverified", "codes": [], "aggregate": False}
    code = candidates[0]
    if old.get(code) != current[code]:
        return {"status": "changed", "codes": [], "aggregate": False}
    if code in available:
        return {"status": "available", "codes": [code], "aggregate": False}
    # Ordinary cities have no standalone SGIS polygon. Require the official
    # parent and every official child to be unchanged and have source geometry.
    children = [c for c, v in current.items() if len(c) == 5 and c != code
                and c.startswith(code[:4]) and v["province"] == row["provinceName"]
                and v["name"].startswith(row["districtName"] + " ")]
    if code.endswith("0") and children and all(c in available and old.get(c) == current[c] for c in children):
        return {"status": "available", "codes": sorted(children), "aggregate": True}
    return {"status": "missing", "codes": [], "aggregate": False}


def polygon_counts(g):
    polygons = list(g.geoms) if g.geom_type == "MultiPolygon" else [g]
    return len(polygons), sum(len(p.interiors) for p in polygons)


def build(sources):
    for name, digest in PINS.items():
        if digest_file(sources / name) != digest:
            raise ValueError(f"Unreviewed source: {name}")
    old, current = (read_links(sources / f"kssc-{date}.xlsx") for date in ("20250401", "20260710"))
    catalogue = json.loads((ROOT / "apps/web/data/region-catalogue.json").read_text())
    forward = pyproj.Transformer.from_crs(5179, 4326, always_xy=True, allow_ballpark=False)
    inverse = pyproj.Transformer.from_crs(4326, 5179, always_xy=True, allow_ballpark=False)
    converted, checks = {}, {}
    with zipfile.ZipFile(sources / "sgis-2025.zip") as archive:
        for level, expected in (("sido", 17), ("sigungu", 252)):
            paths = {ext: next(n for n in archive.namelist() if n.endswith(f"bnd_{level}_00_2025_2Q.{ext}"))
                     for ext in ("shp", "shx", "dbf", "prj")}
            assert pyproj.CRS.from_wkt(archive.read(paths["prj"]).decode()).to_epsg() == 5179
            reader = shapefile.Reader(**{ext: io.BytesIO(archive.read(paths[ext])) for ext in ("shp", "shx", "dbf")}, encoding="utf-8")
            records = [r.as_dict() for r in reader.records()]
            geometries = [shape(s.__geo_interface__) for s in reader.shapes()]
            assert len(geometries) == expected and all(g.is_valid for g in geometries)
            assert shapely.coverage_is_valid(geometries)
            # Simplify the whole coverage together: adjacent polygons retain
            # common edges; islands and holes are not removed to meet a budget.
            simplified = shapely.coverage_simplify(geometries, 50)
            assert shapely.coverage_is_valid(simplified)
            result, errors = {}, []
            for record, original, reduced in zip(records, geometries, simplified):
                assert record["BASE_DATE"] == "20250630"
                assert reduced.is_valid and polygon_counts(original) == polygon_counts(reduced)
                error = abs(reduced.area - original.area) / original.area
                assert error < .005, (record, error)
                geo = transform(forward.transform, reduced)
                assert geo.is_valid and 124 < geo.bounds[0] < geo.bounds[2] < 133 and 32 < geo.bounds[1] < geo.bounds[3] < 39
                restored = transform(inverse.transform, geo)
                assert reduced.hausdorff_distance(restored) < .01
                # Eight decimals retain millimetre-scale vertices while
                # avoiding unnecessary JSON digits. Recheck tiny polygons.
                geo = shape(json.loads(json.dumps(mapping(geo)), parse_float=lambda v: round(float(v), 8)))
                assert geo.is_valid and polygon_counts(geo) == polygon_counts(original)
                result[record["SIDO_CD" if level == "sido" else "SIGUNGU_CD"]] = geo
                errors.append(error)
            converted[level] = result
            assert shapely.coverage_is_valid(list(result.values()))
            checks[level] = {"features": expected, "valid": True, "sharedEdgesValid": True,
                             "polygonAndHoleCountsPreserved": True, "maxRelativeAreaError": max(errors),
                             "projectionRoundTripMetresBelow": .01}
            print(f"{level}: validated {expected} shapes", flush=True)
    folder = ROOT / "apps/web/public/data/boundaries" / VERSION
    folder.mkdir(parents=True, exist_ok=True)
    regions, provinces, national, files = {}, {}, [], {}
    def feature(id_, name, geo, aggregate=False):
        return {"type": "Feature", "properties": {"id": id_, "name": name, "aggregate": aggregate}, "geometry": mapping(geo)}
    def emit(name, features):
        payload = {"type": "FeatureCollection", "version": VERSION, "features": features}
        data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()
        assert len(data) < 4_000_000, (name, len(data))
        (folder / name).write_bytes(data)
        files[name] = {"bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
    for row in catalogue["rows"]:
        link = link_region(row, old, current, converted["sigungu"])
        regions[row["provinceCode"] + ":" + row["districtCode"]] = {**row, **link}
    for province, name in sorted({(r["provinceCode"], r["provinceName"]) for r in catalogue["rows"]}):
        legal = "3600000000" if province == "36110" else province + "00000000"
        candidates = [c for c, v in current.items() if len(c) == 2 and v["province"] == name and v["legal"] == legal]
        code = candidates[0] if len(candidates) == 1 else None
        # Sejong's table encodes only its district-level entry. It has a unique
        # official classification prefix 29 backed by the checked district.
        if province == "36110" and regions["36110:36110"]["status"] == "available":
            code = regions["36110:36110"]["codes"][0][:2]
        available = bool(code and code in converted["sido"] and
                         (old.get(code) == current.get(code) or province == "36110"))
        provinces[province] = {"status": "available" if available else "changed", "code": code, "name": name}
        if available:
            g = converted["sido"][code]
            provinces[province]["bounds"] = list(g.bounds)
            national.append(feature(province, name, g))
        local = []
        for r in regions.values():
            if r["provinceCode"] != province or r["status"] != "available":
                continue
            g = unary_union([converted["sigungu"][c] for c in r["codes"]])
            assert g.is_valid
            r["bounds"] = list(g.bounds)
            local.append(feature(r["districtCode"], r["districtName"], g, r["aggregate"]))
        if local:
            emit(f"province-{province}.json", local)
    emit("national.json", national)
    index = {"version": VERSION, "source": SOURCE, "boundaryDate": "2025-06-30", "crosswalkDate": "2026-07-10",
             "catalogueCollectedAt": catalogue["collectedAt"], "provinces": provinces, "regions": regions}
    (ROOT / "apps/web/data/region-boundaries.json").write_text(json.dumps(index, ensure_ascii=False, indent=2) + "\n")
    report = {"version": VERSION, "sourcePins": PINS, "coordinateDecimals": 8, "software": {"pyproj": pyproj.__version__, "shapely": shapely.__version__, "pyshp": shapefile.__version__, "openpyxl": openpyxl.__version__},
              "simplificationToleranceMetres": 50, "geometry": checks, "counts": dict(Counter(r["status"] for r in regions.values())),
              "aggregateRegions": sum(r["aggregate"] for r in regions.values()), "provinceBoundaries": len(national), "files": files,
              "limitation": "Dated 2025 reference geometry; not certification of current legal boundaries or spatial allocation of statistics"}
    (ROOT / "docs/validation/evidence/2026-09-09-boundary-build.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: v for k, v in report.items() if k not in ("files", "sourcePins")}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sources", required=True, type=Path)
    build(parser.parse_args().sources)
